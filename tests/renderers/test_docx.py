"""DOCX renderer tests.

Unlike the WeasyPrint PDF test, these run everywhere: docxtpl is pure Python
apart from lxml, which ships self-contained wheels.
"""

from __future__ import annotations

import dataclasses
import io

import pytest

from adapters.renderers.docx import DocxRenderer
from domain.errors import RenderError, TemplateNotFoundError
from domain.models import OutputFormat

_DATA = {"community_name": "Communauté ACME", "representative": "Alice Dupont"}


def _render(bundle, data=None, fmt=OutputFormat.DOCX):
    return DocxRenderer().render(bundle, _DATA if data is None else data, fmt, locale="fr-BE")


def _text(content: bytes) -> str:
    from docx import Document

    return "\n".join(p.text for p in Document(io.BytesIO(content)).paragraphs)


def test_render_substitutes_placeholders(docx_bundle):
    artifact = _render(docx_bundle)

    assert artifact.format is OutputFormat.DOCX
    assert artifact.filename == "agreement.docx"
    assert (
        artifact.content_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    body = _text(artifact.content)
    assert "Communauté ACME" in body
    assert "Alice Dupont" in body
    assert "{{" not in body


def test_xml_special_characters_survive(docx_bundle):
    """A .docx body is XML: an unescaped '&' produces a package Word refuses."""
    artifact = _render(
        docx_bundle,
        data={"community_name": "Smith & Co <SRL>", "representative": 'A "B"'},
    )
    body = _text(artifact.content)  # parses the package — corrupt XML raises here
    assert "Smith & Co <SRL>" in body
    assert 'A "B"' in body


def test_missing_field_is_a_hard_error(docx_bundle):
    """StrictUndefined: an official document must never render a silent blank."""
    with pytest.raises(RenderError):
        _render(docx_bundle, data={"community_name": "Only one"})


def test_a_null_value_renders_blank_not_the_word_none(docx_bundle):
    """In a contract, a literal "None" reads as a filled-in value."""
    artifact = _render(docx_bundle, data={"community_name": "ACME", "representative": None})
    body = _text(artifact.content)
    assert "None" not in body
    assert "ACME" in body


def test_render_bytes_are_deterministic(docx_bundle):
    # python-docx lets zipfile stamp members with time.localtime(); the OOXML
    # repack is what keeps the sha256 a function of (template, data).
    assert _render(docx_bundle).content == _render(docx_bundle).content


def test_author_metadata_is_normalised(docx_bundle):
    from docx import Document

    props = Document(io.BytesIO(_render(docx_bundle).content)).core_properties
    assert props.author == "document-generation"
    assert props.last_modified_by == "document-generation"


def test_entrypoint_traversal_is_rejected(docx_bundle):
    evil = dataclasses.replace(
        docx_bundle,
        manifest=docx_bundle.manifest.model_copy(update={"entrypoint": "../escape.docx"}),
    )
    with pytest.raises(TemplateNotFoundError):
        DocxRenderer().render(evil, _DATA, OutputFormat.DOCX, locale=None)


def test_wrong_format_raises_render_error(docx_bundle):
    with pytest.raises(RenderError):
        _render(docx_bundle, fmt=OutputFormat.PDF)
